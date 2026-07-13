"""
tally_report.py -- cached SQL reporting over the Tally HTTP-XML gateway (READ-ONLY).

WHY: fetching balances live for a big multi-year book is slow and can hang Tally's single-threaded
gateway. This pulls the Voucher Register for a date range ONCE into a local SQLite cache, then answers
any number of questions instantly as SQL. Refreshing = re-run with --refresh. Reads never change data.
(Adapted from puneetkeshav/tally-integration's tally_report.py; stdlib-only, cache lives in the
skill's session dir.)

USAGE
  python tally_report.py --list
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --report trial-balance
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --report ledger --arg "HDFC Bank"
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --report pnl
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --report bank-recon \
      --arg "HDFC Bank" --statement stmt.csv          # csv cols: date, debit, credit
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --report tds-vs-26as --form26as 26as.csv
  python tally_report.py --company "My Co" --from 20250401 --to 20260331 --sql "SELECT vtype,COUNT(*) c FROM vouchers GROUP BY vtype"
  python tally_report.py ... --out report.xlsx        # or .csv; console table otherwise
  python tally_report.py ... --refresh                # force a fresh pull from Tally

CACHE TABLES (for --sql):
  vouchers(id, date, vtype, remoteid, narration)
  entries(voucher_id, ledger, dr, cr)          -- dr/cr positive; one row per ledger leg
  inventory(voucher_id, stockitem, in_out, qty, amount)
  masters(name, parent)  /  groups(name, parent)

Deps: none (stdlib); openpyxl only for --out *.xlsx. Company must be loaded; gateway on 9000.
"""
import sys, os, re, sqlite3, csv, argparse
import urllib.request

URL = os.environ.get("TALLY_URL", "http://localhost:9000")
STATE_DIR = os.environ.get("TALLY_SESSION_DIR") or os.path.join(os.getcwd(), "tally-session")
CACHE_DIR = os.path.join(STATE_DIR, "_cache")

# ---- built-in SQL reports: name -> SQL. Use :arg for a parameter (pass via --arg). ----------------
BUILTIN = {
    "trial-balance":   "SELECT ledger, ROUND(SUM(dr),2) dr, ROUND(SUM(cr),2) cr, "
                       "ROUND(SUM(dr)-SUM(cr),2) net FROM entries GROUP BY ledger ORDER BY ledger",
    "group-summary":   "SELECT COALESCE(m.parent,'(unmapped)') grp, ROUND(SUM(e.dr)-SUM(e.cr),2) net "
                       "FROM entries e LEFT JOIN masters m ON m.name=e.ledger GROUP BY grp ORDER BY grp",
    "voucher-summary": "SELECT v.vtype, COUNT(DISTINCT v.id) vouchers, ROUND(SUM(e.dr),2) total "
                       "FROM vouchers v JOIN entries e ON e.voucher_id=v.id GROUP BY v.vtype ORDER BY vouchers DESC",
    "daybook":         "SELECT date, vtype, remoteid, narration FROM vouchers ORDER BY date, id",
    "tds":             "SELECT ledger, ROUND(SUM(dr)-SUM(cr),2) net FROM entries "
                       "WHERE UPPER(ledger) LIKE '%TDS%' GROUP BY ledger ORDER BY ledger",
    "stock":           "SELECT stockitem, ROUND(SUM(CASE WHEN in_out='in' THEN qty END),3) in_qty, "
                       "ROUND(SUM(CASE WHEN in_out='out' THEN qty END),3) out_qty, "
                       "ROUND(SUM(CASE WHEN in_out='in' THEN qty ELSE -qty END),3) net_qty "
                       "FROM inventory GROUP BY stockitem ORDER BY stockitem",
    # parameterised: one ledger's statement (running balance added in Python)
    "ledger":          "SELECT v.date, v.vtype, v.narration, e.dr, e.cr FROM entries e "
                       "JOIN vouchers v ON v.id=e.voucher_id WHERE e.ledger=:arg ORDER BY v.date, v.id",
}

# --------------------------------------------------------------------- pull + cache
def _post(xml, timeout):
    req = urllib.request.Request(URL, data=xml.encode("utf-8"),
                                 headers={"Content-Type": "text/xml"}, method="POST")
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read().decode("utf-8", errors="replace")

def _voucher_register(company, frm, to, timeout):
    xml = ('<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST><TYPE>Data</TYPE>'
           '<ID>Voucher Register</ID></HEADER><BODY><DESC><STATICVARIABLES>'
           '<SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>'
           f'<SVCURRENTCOMPANY>{company}</SVCURRENTCOMPANY>'
           f'<SVFROMDATE>{frm}</SVFROMDATE><SVTODATE>{to}</SVTODATE>'
           '</STATICVARIABLES></DESC></BODY></ENVELOPE>')
    return _post(xml, timeout)

def _list_of_accounts(company, timeout):
    # masters export: names + parents only, no balances -> safe, never hangs (verified)
    xml = ('<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST><TYPE>Data</TYPE>'
           '<ID>List of Accounts</ID></HEADER><BODY><DESC><STATICVARIABLES>'
           '<SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>'
           f'<SVCURRENTCOMPANY>{company}</SVCURRENTCOMPANY>'
           '</STATICVARIABLES></DESC></BODY></ENVELOPE>')
    return _post(xml, timeout)

def _qty(s):
    m = re.search(r"-?[\d.]+", s or "")
    return float(m.group(0)) if m else 0.0

def _unesc(s):
    return (s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            .replace("&quot;", '"').replace("&apos;", "'").replace("&#39;", "'"))

def build_cache(company, frm, to, db_path, timeout):
    vr = _voucher_register(company, frm, to, timeout)
    if len(vr) < 200:
        raise SystemExit(f"Empty/short Voucher Register ({len(vr)} bytes) -- is '{company}' loaded and the range valid?")
    if os.path.exists(db_path):
        os.remove(db_path)
    db = sqlite3.connect(db_path)
    db.executescript("""
      CREATE TABLE vouchers(id INTEGER PRIMARY KEY, date TEXT, vtype TEXT, remoteid TEXT, narration TEXT);
      CREATE TABLE entries(voucher_id INT, ledger TEXT, dr REAL, cr REAL);
      CREATE TABLE inventory(voucher_id INT, stockitem TEXT, in_out TEXT, qty REAL, amount REAL);
      CREATE TABLE masters(name TEXT, parent TEXT);
      CREATE TABLE groups(name TEXT, parent TEXT);
    """)
    vid = 0
    leg_re = re.compile(r"<(ALLLEDGERENTRIES\.LIST|LEDGERENTRIES\.LIST|ACCOUNTINGALLOCATIONS\.LIST)>(.*?)</\1>", re.S)
    for block in re.split(r"(?=<VOUCHER[ >])", vr):
        if not block.startswith("<VOUCHER"):
            continue
        # keep only in-range dates (the VR date filter is sometimes ignored)
        d = (re.search(r"<DATE>(\d{8})</DATE>", block) or [None, ""])[1]
        if not d or d < frm or d > to:
            continue
        vid += 1
        vtype = (re.search(r'VCHTYPE="([^"]*)"', block) or [None, ""])[1]
        rid = (re.search(r'REMOTEID="([^"]*)"', block) or [None, ""])[1]
        nar = (re.search(r"<NARRATION>(.*?)</NARRATION>", block) or [None, ""])[1]
        db.execute("INSERT INTO vouchers VALUES(?,?,?,?,?)", (vid, d, vtype, rid, _unesc(nar)))
        for _tag, seg in leg_re.findall(block):
            ln = re.search(r"<LEDGERNAME>(.*?)</LEDGERNAME>", seg)
            dp = re.search(r"<ISDEEMEDPOSITIVE>(.*?)</ISDEEMEDPOSITIVE>", seg)
            am = re.search(r"<AMOUNT>(-?[\d.]+)</AMOUNT>", seg)
            if ln and dp and am:
                amt = abs(float(am.group(1)))
                isdr = dp.group(1).strip() == "Yes"
                db.execute("INSERT INTO entries VALUES(?,?,?,?)",
                           (vid, _unesc(ln.group(1).strip()), amt if isdr else 0.0, 0.0 if isdr else amt))
        for inv in re.findall(r"<ALLINVENTORYENTRIES\.LIST>(.*?)</ALLINVENTORYENTRIES\.LIST>", block, re.S):
            it = re.search(r"<STOCKITEMNAME>(.*?)</STOCKITEMNAME>", inv)
            dp = re.search(r"<ISDEEMEDPOSITIVE>(.*?)</ISDEEMEDPOSITIVE>", inv)
            q = re.search(r"<ACTUALQTY>(.*?)</ACTUALQTY>", inv)
            am = re.search(r"<AMOUNT>(-?[\d.]+)</AMOUNT>", inv)
            if it and dp:
                db.execute("INSERT INTO inventory VALUES(?,?,?,?,?)",
                           (vid, _unesc(it.group(1).strip()), "in" if dp.group(1).strip() == "Yes" else "out",
                            abs(_qty(q.group(1) if q else "")), abs(float(am.group(1))) if am else 0.0))
    # masters + groups from the safe List of Accounts export (one call, no balances)
    try:
        mt = _list_of_accounts(company, min(timeout, 120))
        for m in re.finditer(r'<LEDGER NAME="([^"]*)"[^>]*>(.*?)</LEDGER>', mt, re.S):
            par = re.search(r"<PARENT[^>]*>(.*?)</PARENT>", m.group(2))
            db.execute("INSERT INTO masters VALUES(?,?)", (_unesc(m.group(1)), _unesc(par.group(1)) if par else ""))
        for m in re.finditer(r'<GROUP NAME="([^"]*)"[^>]*>(.*?)</GROUP>', mt, re.S):
            par = re.search(r"<PARENT[^>]*>(.*?)</PARENT>", m.group(2))
            db.execute("INSERT INTO groups VALUES(?,?)", (_unesc(m.group(1)), _unesc(par.group(1)) if par else ""))
    except Exception as e:
        print(f"  (masters not cached: {type(e).__name__}; group-summary/pnl may be incomplete)", file=sys.stderr)
    db.commit()
    n = db.execute("SELECT COUNT(*) FROM vouchers").fetchone()[0]
    print(f"  cached {n} vouchers, {db.execute('SELECT COUNT(*) FROM entries').fetchone()[0]} legs -> {db_path}")
    return db

# --------------------------------------------------------------------- run + output
def run(db, sql, arg=None, add_running=False):
    cur = db.execute(sql, {"arg": arg} if ":arg" in sql else {})
    cols = [c[0] for c in cur.description]
    rows = [list(r) for r in cur.fetchall()]
    if add_running and {"dr", "cr"} <= set(cols):
        i_dr, i_cr = cols.index("dr"), cols.index("cr")
        cols.append("balance"); bal = 0.0
        for r in rows:
            bal += (r[i_dr] or 0) - (r[i_cr] or 0)
            r.append(round(bal, 2))
    return cols, rows

def output(cols, rows, out):
    if not out:
        w = [max(len(str(c)), *([len(str(r[i])) for r in rows] or [0])) for i, c in enumerate(cols)]
        line = lambda vals: "  ".join(str(v).ljust(w[i]) for i, v in enumerate(vals))
        print(line(cols)); print("  ".join("-" * x for x in w))
        for r in rows:
            print(line(r))
        print(f"({len(rows)} rows)")
    elif out.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.append(cols)
        for r in rows:
            ws.append(r)
        wb.save(out); print(f"wrote {out} ({len(rows)} rows)")
    else:
        with open(out, "w", newline="", encoding="utf-8") as f:
            wr = csv.writer(f); wr.writerow(cols); wr.writerows(rows)
        print(f"wrote {out} ({len(rows)} rows)")

# --------------------------------------------------- computed reports (need logic / CSV inputs)
INCOME_GROUPS = {"Sales Accounts", "Direct Incomes", "Indirect Incomes"}
EXPENSE_GROUPS = {"Purchase Accounts", "Direct Expenses", "Indirect Expenses"}

def _norm_date_any(s):
    s = str(s or "").strip()
    if re.fullmatch(r"\d{8}", s):
        return s
    m = re.match(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})$", s)       # DD-MM-YYYY
    if m:
        return f"{m.group(3)}{int(m.group(2)):02d}{int(m.group(1)):02d}"
    m = re.match(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})$", s)       # YYYY-MM-DD
    if m:
        return f"{m.group(1)}{int(m.group(2)):02d}{int(m.group(3)):02d}"
    return s

def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [{(k or "").strip().lower(): v for k, v in row.items()} for row in csv.DictReader(f)]

def _num(v):
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except (ValueError, AttributeError):
        return 0.0

def rpt_pnl(db, a):
    """Income vs expense by primary P&L group (walks the group tree), signed, + net profit."""
    grp = dict(db.execute("SELECT name, parent FROM groups"))
    led = dict(db.execute("SELECT name, parent FROM masters"))
    def primary(ledger):
        node, seen = led.get(ledger), set()
        while node and node not in seen:
            if node in INCOME_GROUPS or node in EXPENSE_GROUPS:
                return node
            seen.add(node); node = grp.get(node)
        return None
    bygroup, income, expense = {}, 0.0, 0.0
    for ledger, dr, cr in db.execute("SELECT ledger, SUM(dr), SUM(cr) FROM entries GROUP BY ledger"):
        p = primary(ledger)
        if not p:
            continue
        amt = (cr - dr) if p in INCOME_GROUPS else (dr - cr)
        bygroup[p] = bygroup.get(p, 0) + amt
        income += amt if p in INCOME_GROUPS else 0
        expense += amt if p in EXPENSE_GROUPS else 0
    rows = [[g, round(bygroup[g], 2)] for g in ("Sales Accounts", "Direct Incomes", "Indirect Incomes") if g in bygroup]
    rows.append(["-- Total Income --", round(income, 2)])
    rows += [[g, round(bygroup[g], 2)] for g in ("Purchase Accounts", "Direct Expenses", "Indirect Expenses") if g in bygroup]
    rows.append(["-- Total Expense --", round(expense, 2)])
    rows.append(["== NET PROFIT ==", round(income - expense, 2)])
    return ["section", "amount"], rows

def rpt_bank_recon(db, a):
    """Diff a bank ledger's book movement vs a statement CSV by NET DAILY movement (owners
    consolidate same-day entries, so per-line matching double-counts)."""
    if not (a.arg and a.statement):
        raise SystemExit('bank-recon needs --arg "<bank ledger>" and --statement <csv with date,debit,credit>')
    book = {}
    for d, dr, cr in db.execute("SELECT v.date, SUM(e.dr), SUM(e.cr) FROM entries e "
                                "JOIN vouchers v ON v.id=e.voucher_id WHERE e.ledger=:arg GROUP BY v.date",
                                {"arg": a.arg}):
        book[d] = round((dr or 0) - (cr or 0), 2)                  # book net = Dr - Cr (money in +)
    stmt = {}
    for row in _read_csv(a.statement):
        d = _norm_date_any(row.get("date"))
        stmt[d] = round(stmt.get(d, 0) + _num(row.get("credit")) - _num(row.get("debit")), 2)  # money in +
    rows, tb, ts = [], 0.0, 0.0
    for d in sorted(set(book) | set(stmt)):
        b, s = book.get(d, 0), stmt.get(d, 0); tb += b; ts += s
        if abs(b - s) >= 0.005:
            rows.append([d, b, s, round(b - s, 2)])
    rows.append(["TOTAL", round(tb, 2), round(ts, 2), round(tb - ts, 2)])
    return ["date", "book_net", "stmt_net", "diff(only mismatched dates)"], rows

def rpt_tds_26as(db, a):
    """Book TDS-credit ledgers vs a 26AS CSV (Active rows only; Inactive/duplicate rows excluded)."""
    if not a.form26as:
        raise SystemExit("tds-vs-26as needs --form26as <26AS csv> (cols: Section, TDS Deducted, Status)")
    book = db.execute("SELECT ROUND(SUM(dr)-SUM(cr),2) FROM entries "
                      "WHERE UPPER(ledger) LIKE '%TDS%' AND UPPER(ledger) NOT LIKE '%PAYABLE%'").fetchone()[0] or 0
    bysec, active, inactive = {}, 0.0, 0.0
    for row in _read_csv(a.form26as):
        tds = _num(row.get("tds deducted"))
        if (row.get("status") or "active").strip().lower() == "active":
            sec = (row.get("section") or "?").strip()
            bysec[sec] = bysec.get(sec, 0) + tds; active += tds
        else:
            inactive += tds
    rows = [[f"26AS TDS {s} (Active)", round(bysec[s], 2)] for s in sorted(bysec)]
    rows += [["26AS TDS total (Active)", round(active, 2)],
             ["Book TDS ledgers (net, excl. Payable)", round(book, 2)],
             ["Difference (book - 26AS)", round(book - active, 2)]]
    if inactive:
        rows.append(["(excluded: 26AS 'Inactive' rows)", round(inactive, 2)])
    return ["line", "amount"], rows

COMPUTED = {
    "pnl":         (rpt_pnl,        "P&L: income vs expense by primary group + net profit"),
    "bank-recon":  (rpt_bank_recon, "diff a bank ledger vs a statement CSV by net daily movement (--arg <ledger> --statement <csv>)"),
    "tds-vs-26as": (rpt_tds_26as,   "book TDS ledgers vs a 26AS CSV (--form26as <csv>)"),
}

def main():
    ap = argparse.ArgumentParser(description="Cached SQL reporting over the Tally gateway (read-only).")
    ap.add_argument("--company"); ap.add_argument("--from", dest="frm"); ap.add_argument("--to")
    ap.add_argument("--report"); ap.add_argument("--sql"); ap.add_argument("--arg")
    ap.add_argument("--statement", help="statement CSV (for bank-recon)")
    ap.add_argument("--form26as", help="26AS CSV (for tds-vs-26as)")
    ap.add_argument("--out"); ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--timeout", type=int, default=300, help="seconds; heavy first pulls need minutes, not 60s")
    a = ap.parse_args()
    if a.list:
        print("SQL reports:")
        for k in sorted(BUILTIN):
            print(f"  {k}{'   (needs --arg)' if ':arg' in BUILTIN[k] else ''}")
        print("Computed reports:")
        for k in sorted(COMPUTED):
            print(f"  {k}   -- {COMPUTED[k][1]}")
        return
    if not (a.company and a.frm and a.to):
        ap.error("--company, --from, --to are required (or use --list)")
    os.makedirs(CACHE_DIR, exist_ok=True)
    db_path = os.path.join(CACHE_DIR, f"{re.sub(r'[^A-Za-z0-9]', '_', a.company)}_{a.frm}_{a.to}.sqlite")
    if a.refresh or not os.path.exists(db_path):
        print(f"Building cache for {a.company!r} {a.frm}-{a.to} (first pull can be slow on a big book)...")
        build_cache(a.company, a.frm, a.to, db_path, a.timeout).close()
    db = sqlite3.connect(db_path)
    if a.sql:
        output(*run(db, a.sql), a.out)
    elif a.report in COMPUTED:
        output(*COMPUTED[a.report][0](db, a), a.out)
    elif a.report:
        if a.report not in BUILTIN:
            ap.error(f"unknown report {a.report!r}; use --list")
        sql = BUILTIN[a.report]
        if ":arg" in sql and not a.arg:
            ap.error(f"report {a.report!r} needs --arg")
        output(*run(db, sql, a.arg, add_running=(a.report == "ledger")), a.out)
    else:
        ap.error("give --report NAME or --sql \"...\" (or --list)")

if __name__ == "__main__":
    main()
